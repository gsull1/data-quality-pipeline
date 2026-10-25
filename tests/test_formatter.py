import pandas as pd
from openpyxl import Workbook
from src.formatter import create_quality_report_workbook

def test_create_quality_report_workbook():
    df_cont = pd.DataFrame({"Feature": ["col1"], "Mean": [10.5]})
    df_cat = pd.DataFrame({"Feature": ["col2"], "Top Mode": ["A"]})

    wb = create_quality_report_workbook(df_cont, df_cat)

    assert isinstance(wb, Workbook)
    assert "Continuous Profiling" in wb.sheetnames
    assert "Categorical Profiling" in wb.sheetnames

    ws_cont = wb["Continuous Profiling"]
    assert ws_cont.cell(row=1,column=1).value == "Feature"
    assert ws_cont.cell(row=2,column=1).value == "col1"
    