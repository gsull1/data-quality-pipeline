"""Excel workbook formatting engine using openpyxl."""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def create_quality_report_workbook(
    df_continuous: pd.DataFrame,
    df_categorical: pd.DataFrame
) -> Workbook:
    """Generates a styled openpyxl Workbook containing Continuous and Categorical quality sheets."""
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Style Definitions
    header_font = Font(name="Calibri", size = 11, bold = True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    border_thin = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    align_center = Alignment(horizontal="center",vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    sheets_data = [
        ("Continuous Profiling", df_continuous),
        ("Categorical Profiling", df_categorical)
    ]

    for sheet_name, df_data in sheets_data:
        ws = wb.create_sheet(title=sheet_name)

        if df_data is None or df_data.empty:
            ws.append(["No features available for this data type."])
            continue

        #Write dataframe to sheet
        for r_idx, row in enumerate(dataframe_to_rows(df_data, index=False, header = True), 1):
            ws.append(row)
            for c_idx in range(1, len(row)):
                cell = ws.cell(row=r_idx, column = c_idx)
                cell.border = cell_border

                # Style Header Row vs Data Rows
                if r_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = align_center
                else:
                    if c_idx == 1:
                        cell.alignment = align_left
                    else:
                        cell.alignment = align_center

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    return wb